"""Excel styling constants shared across every assessment report tab."""
from openpyxl.styles import Border, Font, PatternFill, Side

# Excel styling
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FONT = Font(bold=True, size=12)
TITLE_FONT = Font(bold=True, size=14)
THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# Status colors
STATUS_COLORS = {
    'protected': PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    'partial': PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    'unprotected': PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    'info': PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid"),
}

# Provider colors
PROVIDER_COLORS = {
    'AWS': PatternFill(start_color="FF9900", end_color="FF9900", fill_type="solid"),
    'Azure': PatternFill(start_color="0078D4", end_color="0078D4", fill_type="solid"),
    'GCP': PatternFill(start_color="4285F4", end_color="4285F4", fill_type="solid"),
    'M365': PatternFill(start_color="D83B01", end_color="D83B01", fill_type="solid"),
}
