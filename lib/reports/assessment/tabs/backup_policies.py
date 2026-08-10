"""Backup Policies tab: AWS Backup plans, DLM policies, Azure policies (included/excluded scope)."""
from collections import defaultdict
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from ..excel_helpers import (
    set_column_widths,
    write_data_row,
    write_header_row,
    write_section_header,
)
from ..styles import (
    STATUS_COLORS,
    TITLE_FONT,
)


def generate_backup_policies(wb: Workbook, resources: List[Dict]) -> None:
    """Generate Backup Policies tab.

    Shows each AWS Backup plan (with its selections, included/excluded resources
    and retention rules), AWS DLM lifecycle policies (target/exclude tags,
    schedules, retention), and Azure Backup policies (vault, type, retention).

    The apples-to-apples note at the bottom explains that Option 2 of the
    Sizing Inputs tab reflects exactly the resources covered by these policies.
    """
    ws = wb.create_sheet(title="Backup Policies")

    row = 1
    ws.cell(row=row, column=1, value="Backup Policies").font = TITLE_FONT
    ws.cell(row=row + 1, column=1,
            value="Included/excluded resource scope for AWS Backup, DLM, and Azure Backup policies")
    row += 3

    # -------------------------------------------------------------------------
    # SECTION 1: AWS Backup Plans
    # -------------------------------------------------------------------------
    aws_plans = [r for r in resources if r.get('resource_type') == 'aws:backup:plan']
    aws_selections = [r for r in resources if r.get('resource_type') == 'aws:backup:selection']

    # Build selection lookup keyed by (plan_id, account_id, region). Org-managed
    # AWS Backup plans (orgs/...) are deployed to every account in the org, so the
    # same plan_id appears in each account; selections are per-account and must
    # not fan out across every account that shares the plan_id.
    selections_by_plan: Dict[tuple, List[Dict]] = defaultdict(list)
    for sel in aws_selections:
        plan_id = (sel.get('metadata') or {}).get('backup_plan_id', '')
        if plan_id:
            key = (plan_id, sel.get('account_id', ''), sel.get('region', ''))
            selections_by_plan[key].append(sel)

    row = write_section_header(ws, row, "AWS Backup Plans",
                               f"({len(aws_plans)} plans found)")

    if not aws_plans:
        ws.cell(row=row, column=1, value="No AWS Backup plans found")
        row += 2
    else:
        for plan in sorted(aws_plans, key=lambda r: r.get('name', '')):
            meta = plan.get('metadata') or {}
            plan_id = meta.get('backup_plan_id', plan.get('resource_id', ''))
            plan_name = plan.get('name', plan_id)
            last_run = meta.get('last_execution_date') or 'N/A'

            # Plan header
            ws.cell(row=row, column=1, value=f"Plan: {plan_name}").font = Font(bold=True)
            ws.cell(row=row, column=2, value=f"Account: {plan.get('account_id', '')}")
            ws.cell(row=row, column=3, value=f"Region: {plan.get('region', '')}")
            ws.cell(row=row, column=4, value=f"Last Run: {last_run}")
            row += 1

            # Rules (schedule + retention)
            rules = meta.get('rules') or []
            if rules:
                write_header_row(ws, row, ["  Rule Name", "Schedule", "Retention (days)",
                                           "Cold Storage After (days)", "Target Vault"])
                row += 1
                for rule in rules:
                    delete_after = rule.get('delete_after_days')
                    cold_after = rule.get('move_to_cold_after_days')
                    write_data_row(ws, row, [
                        f"  {rule.get('rule_name', '')}",
                        rule.get('schedule', ''),
                        delete_after if delete_after is not None else 'Indefinite',
                        cold_after if cold_after is not None else '—',
                        rule.get('target_vault', '')
                    ])
                    row += 1
                row += 1

            # Selections (included/excluded resources), scoped to this plan instance's account/region
            plan_sels = selections_by_plan.get(
                (plan_id, plan.get('account_id', ''), plan.get('region', '')), []
            )
            if plan_sels:
                write_header_row(ws, row, ["  Selection Name", "Included Resources / Tags",
                                           "Excluded Resources", "IAM Role"])
                row += 1
                for sel in plan_sels:
                    smeta = sel.get('metadata') or {}
                    included_arns = smeta.get('resources') or []
                    excluded_arns = smeta.get('not_resources') or []
                    tag_filters = smeta.get('list_of_tags') or []

                    # Build a human-readable inclusion string
                    if included_arns:
                        included_str = '; '.join(str(a) for a in included_arns[:5])
                        if len(included_arns) > 5:
                            included_str += f' (+{len(included_arns) - 5} more)'
                    elif tag_filters:
                        tag_parts = []
                        for tf in tag_filters[:4]:
                            if isinstance(tf, dict):
                                tag_parts.append(f"{tf.get('ConditionKey', '')}={tf.get('ConditionValue', '')}")
                        included_str = 'Tags: ' + ', '.join(tag_parts)
                        if len(tag_filters) > 4:
                            included_str += f' (+{len(tag_filters) - 4} more tags)'
                    else:
                        included_str = 'All resources'

                    excluded_str = '; '.join(str(a) for a in excluded_arns[:3]) if excluded_arns else '—'
                    if len(excluded_arns) > 3:
                        excluded_str += f' (+{len(excluded_arns) - 3} more)'

                    # Extract role name from ARN for readability
                    iam_role = smeta.get('iam_role_arn', '')
                    role_name = iam_role.split('/')[-1] if '/' in iam_role else iam_role

                    write_data_row(ws, row, [
                        f"  {sel.get('name', smeta.get('selection_id', ''))}",
                        included_str,
                        excluded_str,
                        role_name
                    ])
                    row += 1
            else:
                ws.cell(row=row, column=1, value="  (No resource selections found for this plan)")
                row += 1

            row += 1

    row += 1

    # -------------------------------------------------------------------------
    # SECTION 2: AWS DLM (Data Lifecycle Manager) Policies
    # -------------------------------------------------------------------------
    dlm_policies_collected = [r for r in resources if r.get('resource_type') == 'aws:dlm:lifecycle-policy']

    # Also gather any DLM policy IDs inferred from snapshot tags (may predate collection)
    snapshot_dlm_ids: Dict[str, Dict] = {}
    for r in resources:
        if r.get('resource_type') == 'aws:ec2:snapshot':
            tags = r.get('tags') or {}
            pid = tags.get('aws:dlm:lifecycle-policy-id')
            if pid:
                if pid not in snapshot_dlm_ids:
                    snapshot_dlm_ids[pid] = {
                        'count': 0,
                        'size_gb': 0.0,
                        'schedule_name': tags.get('aws:dlm:lifecycle-schedule-name', 'Unknown'),
                        'accounts': set(),
                        'regions': set(),
                    }
                snapshot_dlm_ids[pid]['count'] += 1
                snapshot_dlm_ids[pid]['size_gb'] += r.get('size_gb', 0) or 0
                snapshot_dlm_ids[pid]['accounts'].add(r.get('account_id', ''))
                snapshot_dlm_ids[pid]['regions'].add(r.get('region', ''))

    collected_ids = {(r.get('metadata') or {}).get('policy_id', '') for r in dlm_policies_collected}

    row = write_section_header(ws, row, "AWS DLM Lifecycle Policies",
                               f"({len(dlm_policies_collected)} collected, "
                               f"{len(snapshot_dlm_ids)} observed in snapshots)")

    if not dlm_policies_collected and not snapshot_dlm_ids:
        ws.cell(row=row, column=1, value="No DLM lifecycle policies found")
        row += 2
    else:
        # --- Collected policy definitions ---
        if dlm_policies_collected:
            ws.cell(row=row, column=1, value="Policy Definitions").font = Font(bold=True, italic=True)
            row += 1
            write_header_row(ws, row, ["Policy ID", "Description", "State", "Resource Type",
                                       "Included (Target Tags)", "Excluded (Tags / Boot Vol)",
                                       "Account", "Region"])
            row += 1

            for pol in sorted(dlm_policies_collected, key=lambda r: r.get('name', '')):
                meta = pol.get('metadata') or {}
                pid = meta.get('policy_id', '')
                state = meta.get('state', '')
                state_fill = STATUS_COLORS['protected'] if state == 'ENABLED' else STATUS_COLORS['unprotected']

                target_tags = meta.get('target_tags') or []
                included_str = ', '.join(target_tags[:5]) if target_tags else 'All volumes in region'
                if len(target_tags) > 5:
                    included_str += f' (+{len(target_tags) - 5} more)'

                exclude_parts = []
                if meta.get('exclude_boot_volume'):
                    exclude_parts.append('Boot volumes excluded')
                excl_tags = meta.get('exclude_data_volume_tags') or []
                if excl_tags:
                    exclude_parts.append('Tags: ' + ', '.join(excl_tags[:3]))
                excluded_str = '; '.join(exclude_parts) if exclude_parts else '—'

                write_data_row(ws, row, [
                    pid,
                    meta.get('description', ''),
                    state,
                    meta.get('resource_type', ''),
                    included_str,
                    excluded_str,
                    pol.get('account_id', ''),
                    pol.get('region', '')
                ])
                ws.cell(row=row, column=3).fill = state_fill
                row += 1

            row += 1

            # Schedules sub-table
            ws.cell(row=row, column=1, value="DLM Schedules").font = Font(bold=True, italic=True)
            row += 1
            write_header_row(ws, row, ["Policy ID", "Schedule Name", "Cron / Interval",
                                       "Retain Count", "Retain Period", "Cross-Region Copies"])
            row += 1

            for pol in sorted(dlm_policies_collected, key=lambda r: r.get('name', '')):
                meta = pol.get('metadata') or {}
                pid = meta.get('policy_id', '')
                for sched in (meta.get('schedules') or []):
                    cron = sched.get('cron_expression', '')
                    if not cron and sched.get('interval'):
                        cron = f"Every {sched['interval']} {sched.get('interval_unit', '')}(s)"
                    retain_count = sched.get('retain_count', '')
                    retain_period = ''
                    if sched.get('retain_interval'):
                        retain_period = f"{sched['retain_interval']} {sched.get('retain_interval_unit', '')}"
                    cross_copies = ', '.join(sched.get('cross_region_copies') or []) or '—'
                    write_data_row(ws, row, [
                        pid,
                        sched.get('name', ''),
                        cron,
                        retain_count if retain_count != '' else '—',
                        retain_period or '—',
                        cross_copies
                    ])
                    row += 1

            row += 2

        # --- DLM policies observed only in snapshot tags (no definition collected) ---
        unmatched = {pid: info for pid, info in snapshot_dlm_ids.items() if pid not in collected_ids}
        if unmatched:
            ws.cell(row=row, column=1,
                    value="DLM Policies observed in snapshots (definition not collected)").font = Font(bold=True, italic=True)
            row += 1
            write_header_row(ws, row, ["Policy ID", "Observed Schedule Name",
                                       "Snapshot Count", "Total Size (GB)", "Accounts", "Regions"])
            row += 1
            for pid, info in sorted(unmatched.items(), key=lambda x: x[1]['count'], reverse=True):
                write_data_row(ws, row, [
                    pid,
                    info.get('schedule_name', 'Unknown'),
                    info['count'],
                    round(info['size_gb'], 1),
                    ', '.join(sorted(info['accounts'])),
                    ', '.join(sorted(info['regions']))
                ])
                row += 1
            row += 2

    # -------------------------------------------------------------------------
    # SECTION 3: Azure Backup Policies
    # -------------------------------------------------------------------------
    azure_policies = [r for r in resources if r.get('resource_type') == 'azure:backup:policy']
    azure_protected_items = [r for r in resources if r.get('resource_type') == 'azure:backup:protecteditem']

    # Count protected items per vault for context
    items_per_vault: Dict[str, int] = defaultdict(int)
    for item in azure_protected_items:
        vault = (item.get('metadata') or {}).get('vault_name', '')
        if vault:
            items_per_vault[vault] += 1

    row = write_section_header(ws, row, "Azure Backup Policies",
                               f"({len(azure_policies)} policies across "
                               f"{len({(r.get('metadata') or {}).get('vault_name', '') for r in azure_policies} - {''})} vaults)")

    if not azure_policies:
        ws.cell(row=row, column=1, value="No Azure Backup policies found")
        row += 2
    else:
        write_header_row(ws, row, ["Policy Name", "Vault", "Workload / Type",
                                   "Daily Retention (days)", "Protected Items in Vault",
                                   "Subscription", "Region"])
        row += 1

        workload_order = {'AzureIaasVM': 0, 'AzureWorkload': 1, 'AzureStorage': 2}
        sorted_policies = sorted(
            azure_policies,
            key=lambda r: (
                (r.get('metadata') or {}).get('vault_name', ''),
                workload_order.get((r.get('metadata') or {}).get('policy_type', ''), 9),
                r.get('name', '')
            )
        )

        for pol in sorted_policies:
            meta = pol.get('metadata') or {}
            vault_name = meta.get('vault_name', '')
            policy_type = meta.get('policy_type', 'unknown')
            retention = meta.get('retention_days')
            item_count = items_per_vault.get(vault_name, 0)

            write_data_row(ws, row, [
                pol.get('name', ''),
                vault_name,
                policy_type,
                retention if retention is not None else 'See vault settings',
                item_count if item_count else '—',
                pol.get('subscription_id', ''),
                pol.get('region', '')
            ])
            row += 1

        row += 2

    # -------------------------------------------------------------------------
    # APPLES-TO-APPLES NOTE
    # -------------------------------------------------------------------------
    row = write_section_header(ws, row, "Apples-to-Apples Scope Note")

    note_lines = [
        ("What \"Apples-to-Apples\" means:",
         "Option 2 in the Sizing Inputs tab (Currently Protected Only) is sized against the exact "
         "resources covered by the backup policies listed above. This makes the Cohesity sizing "
         "a direct, like-for-like replacement of your current backup footprint."),
        ("AWS Backup — included scope:",
         "Resources matched by the ARN lists or tag conditions in each Backup Selection above. "
         "Resources in 'not_resources' lists are explicitly excluded from that plan's coverage."),
        ("AWS DLM — included scope:",
         "EBS volumes or EC2 instances that carry the Target Tags listed for each DLM policy. "
         "Volumes tagged with Exclude Data Volume Tags, and boot volumes when "
         "Exclude Boot Volume = True, are not snapshotted by that policy."),
        ("Azure Backup — included scope:",
         "VMs, SQL, and storage items explicitly registered to a Recovery Services vault and "
         "assigned a backup policy. Items in the same vault without a policy assignment, "
         "and resources not registered to any vault, are not covered."),
        ("Why this matters for sizing:",
         "If a resource is in scope for any backup policy it appears in the "
         "'protected' count. Resources outside all policy scopes appear as 'unprotected'. "
         "The Cohesity sizing for Option 2 will match your current protected footprint — "
         "neither over-counting nor under-counting relative to what you back up today."),
    ]

    write_header_row(ws, row, ["Topic", "Explanation"])
    row += 1

    for topic, explanation in note_lines:
        ws.cell(row=row, column=1, value=topic).font = Font(bold=True)
        cell = ws.cell(row=row, column=2, value=explanation)
        cell.alignment = Alignment(wrap_text=True)
        row += 1

    row += 1

    # Column widths
    set_column_widths(ws, {'A': 40, 'B': 60, 'C': 25, 'D': 22, 'E': 22, 'F': 35, 'G': 18, 'H': 18})
    ws.row_dimensions[1].height = 18
    ws.freeze_panes = 'A3'


