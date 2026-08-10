"""Protection Analysis tab: coverage percentages and snapshot analysis."""
from typing import Dict, List

from openpyxl import Workbook

from ..analysis import (
    analyze_protection_status,
    analyze_snapshots,
)
from ..excel_helpers import (
    format_percent,
    set_column_widths,
    write_data_row,
    write_header_row,
    write_section_header,
)
from ..styles import (
    STATUS_COLORS,
    TITLE_FONT,
)


def generate_protection_analysis(wb: Workbook, resources: List[Dict]) -> None:
    """Generate Protection Analysis tab."""
    ws = wb.create_sheet(title="Protection Analysis")

    row = 1

    # Title
    ws.cell(row=row, column=1, value="Protection Analysis").font = TITLE_FONT
    row += 2

    protection = analyze_protection_status(resources)

    # === Coverage Summary ===
    row = write_section_header(ws, row, "Protection Coverage Summary")

    coverage_data = [
        ("Total Protectable Resources", protection['total_protectable'], ""),
        ("Protected Resources", protection['protected_count'],
         format_percent(protection['coverage_percent'])),
        ("Unprotected Resources", protection['unprotected_count'],
         format_percent(100 - protection['coverage_percent'])),
    ]

    write_header_row(ws, row, ["Metric", "Count", "Percentage"])
    row += 1

    for label, count, pct in coverage_data:
        write_data_row(ws, row, [label, count, pct])
        if "Unprotected" in label:
            ws.cell(row=row, column=1).fill = STATUS_COLORS['unprotected']
        elif "Protected" in label and "Total" not in label:
            ws.cell(row=row, column=1).fill = STATUS_COLORS['protected']
        row += 1

    row += 2

    # === Size-based Coverage ===
    row = write_section_header(ws, row, "Size-based Coverage")

    write_header_row(ws, row, ["Status", "Size (GB)", "Size (TB)", "Percentage"])
    row += 1

    total_size = protection['total_size_gb'] or 1  # Avoid division by zero

    for status, size_key, fill in [
        ("Protected", 'protected_size_gb', STATUS_COLORS['protected']),
        ("Unprotected", 'unprotected_size_gb', STATUS_COLORS['unprotected']),
    ]:
        size_gb = protection[size_key]
        pct = (size_gb / total_size * 100) if total_size else 0
        write_data_row(ws, row, [
            status,
            round(size_gb, 1),
            round(size_gb / 1024, 2),
            format_percent(pct)
        ])
        ws.cell(row=row, column=1).fill = fill
        row += 1

    row += 2

    # === Snapshot Analysis ===
    row = write_section_header(ws, row, "Snapshot Inventory")

    snapshot_analysis = analyze_snapshots(resources)

    ws.cell(row=row, column=1, value="Total Snapshots")
    ws.cell(row=row, column=2, value=snapshot_analysis['total_count'])
    row += 1
    ws.cell(row=row, column=1, value="Total Snapshot Size (GB)")
    ws.cell(row=row, column=2, value=round(snapshot_analysis['total_size_gb'], 1))
    row += 1
    ws.cell(row=row, column=1, value="Total Snapshot Size (TB)")
    ws.cell(row=row, column=2, value=round(snapshot_analysis['total_size_gb'] / 1024, 2))
    row += 2

    if snapshot_analysis['by_type']:
        write_header_row(ws, row, ["Snapshot Type", "Count", "Size (GB)"])
        row += 1

        for stype, data in sorted(snapshot_analysis['by_type'].items()):
            write_data_row(ws, row, [stype, data['count'], round(data['size_gb'], 1)])
            row += 1

    # Set column widths
    set_column_widths(ws, {'A': 35, 'B': 15, 'C': 15, 'D': 15})

    ws.freeze_panes = 'A3'


