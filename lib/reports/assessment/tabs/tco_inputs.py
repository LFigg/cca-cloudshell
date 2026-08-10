"""TCO Inputs tab: current backup costs and Cohesity TCO calculator inputs."""
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Font

from ..analysis import (
    categorize_resources,
)
from ..excel_helpers import (
    format_currency,
    set_column_widths,
    write_data_row,
    write_header_row,
    write_section_header,
)
from ..styles import (
    TITLE_FONT,
)


def generate_tco_inputs(wb: Workbook, resources: List[Dict],
                        cost_data: Dict) -> None:
    """Generate TCO Inputs tab."""
    ws = wb.create_sheet(title="TCO Inputs")

    row = 1

    # Title
    ws.cell(row=row, column=1, value="TCO Calculator Inputs").font = TITLE_FONT
    ws.cell(row=row + 1, column=1, value="(Use these values for Cohesity TCO analysis)")
    row += 3

    # === Current Costs Section ===
    row = write_section_header(ws, row, "Current Backup/Storage Costs",
                                "(Monthly costs from cloud billing)")

    if cost_data.get('total_cost', 0) > 0:
        ws.cell(row=row, column=1, value="Total Monthly Cost")
        ws.cell(row=row, column=2, value=format_currency(cost_data['total_cost']))
        row += 2

        # By provider
        if cost_data.get('by_provider'):
            write_header_row(ws, row, ["Provider", "Category", "Monthly Cost"])
            row += 1

            for provider, pdata in sorted(cost_data['by_provider'].items()):
                # Provider total
                write_data_row(ws, row, [
                    provider.upper(),
                    "Total",
                    format_currency(pdata.get('total', 0))
                ])
                ws.cell(row=row, column=1).font = Font(bold=True)
                row += 1

                # Category breakdown
                for category, cost in sorted(pdata.get('categories', {}).items()):
                    write_data_row(ws, row, ["", category, format_currency(cost)])
                    row += 1
    else:
        ws.cell(row=row, column=1, value="No cost data available")
        ws.cell(row=row + 1, column=1, value="Run 'python cost_collect.py' to collect backup cost data")
        row += 3

    row += 2

    # === Sizing Inputs for TCO ===
    row = write_section_header(ws, row, "TCO Sizing Inputs")

    categories = categorize_resources(resources)

    # Calculate totals
    total_size = sum(
        data['size_gb'] for cat, data in categories.items()
        if cat not in ['Snapshots', 'Backup Services', 'Other']
    )

    tco_inputs = [
        ("Total Protected Data (TB)", round(total_size / 1024, 2)),
        ("Estimated Daily Change Rate (%)", "2-5% (typical)"),
        ("Retention Period (days)", "30 (adjust as needed)"),
        ("Number of Replicas", "1-2 (disaster recovery)"),
    ]

    for label, value in tco_inputs:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 2

    # === Projected Annual Costs ===
    row = write_section_header(ws, row, "Projected Annual Costs",
                                "(Based on current monthly costs)")

    if cost_data.get('total_cost', 0) > 0:
        monthly = cost_data['total_cost']
        tco_data = [
            ("Current Monthly Cost", format_currency(monthly)),
            ("Projected Annual Cost", format_currency(monthly * 12)),
            ("Projected 3-Year Cost", format_currency(monthly * 36)),
            ("Projected 5-Year Cost", format_currency(monthly * 60)),
        ]

        for label, value in tco_data:
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=value)
            row += 1

    # Set column widths
    set_column_widths(ws, {'A': 35, 'B': 25, 'C': 20})


