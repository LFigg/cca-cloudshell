"""Account Detail tab: multi-account/subscription breakdown."""
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Font

from ..analysis import (
    analyze_accounts,
)
from ..excel_helpers import (
    set_column_widths,
    write_data_row,
    write_header_row,
    write_section_header,
)
from ..styles import (
    PROVIDER_COLORS,
    TITLE_FONT,
)


def generate_account_detail(wb: Workbook, resources: List[Dict]) -> None:
    """Generate Account Detail tab."""
    ws = wb.create_sheet(title="Account Detail")

    row = 1

    # Title
    ws.cell(row=row, column=1, value="Account/Subscription Detail").font = TITLE_FONT
    row += 2

    accounts = analyze_accounts(resources)

    # === Account Summary ===
    row = write_section_header(ws, row, "Account Summary")

    write_header_row(ws, row, [
        "Account/Subscription", "Provider", "Regions", "Resources", "Size (GB)", "Size (TB)"
    ])
    row += 1

    # Sort by size descending
    sorted_accounts = sorted(
        accounts.items(),
        key=lambda x: x[1]['size_gb'],
        reverse=True
    )

    for account_id, data in sorted_accounts:
        regions = ', '.join(sorted(data['regions']))
        write_data_row(ws, row, [
            account_id,
            data['provider'],
            regions,
            data['count'],
            round(data['size_gb'], 1),
            round(data['size_gb'] / 1024, 2)
        ])

        # Color by provider
        provider = data['provider']
        if provider in PROVIDER_COLORS:
            ws.cell(row=row, column=2).fill = PROVIDER_COLORS[provider]
            ws.cell(row=row, column=2).font = Font(color="FFFFFF")

        row += 1

    row += 2

    # === Resource Type Breakdown by Account ===
    row = write_section_header(ws, row, "Resource Types by Account")

    write_header_row(ws, row, ["Account", "Resource Type", "Count", "Size (GB)"])
    row += 1

    for account_id, data in sorted_accounts:
        for rtype, count in sorted(data['types'].items()):
            # Use pre-computed size from analyze_accounts
            type_size = data.get('type_sizes', {}).get(rtype, 0)
            write_data_row(ws, row, [
                account_id,
                rtype,
                count,
                round(type_size, 1)
            ])
            row += 1

    # Set column widths
    set_column_widths(ws, {'A': 25, 'B': 35, 'C': 15, 'D': 12, 'E': 15, 'F': 12})

    ws.freeze_panes = 'A4'


