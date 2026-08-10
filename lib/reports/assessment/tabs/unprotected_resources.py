"""Unprotected Resources tab: prioritized list for protection planning."""
from typing import Dict, List

from openpyxl import Workbook

from ..analysis import (
    analyze_protection_status,
    get_provider,
)
from ..excel_helpers import (
    set_column_widths,
    write_data_row,
    write_header_row,
    write_section_header,
)
from ..styles import (
    STATUS_COLORS,
    TITLE_FONT,
)


def generate_unprotected_resources(wb: Workbook, resources: List[Dict]) -> None:
    """Generate Unprotected Resources tab."""
    ws = wb.create_sheet(title="Unprotected Resources")

    row = 1

    # Title
    ws.cell(row=row, column=1, value="Unprotected Resources").font = TITLE_FONT
    ws.cell(row=row + 1, column=1, value="(Prioritized list for protection planning)")
    row += 3

    protection = analyze_protection_status(resources)
    unprotected = protection['unprotected_resources']

    # Sort by size descending (prioritize largest)
    unprotected_sorted = sorted(
        unprotected,
        key=lambda x: x.get('size_gb', 0) or 0,
        reverse=True
    )

    # === Summary ===
    row = write_section_header(ws, row, "Summary")

    ws.cell(row=row, column=1, value="Total Unprotected Resources")
    ws.cell(row=row, column=2, value=len(unprotected))
    row += 1
    ws.cell(row=row, column=1, value="Total Unprotected Size (GB)")
    ws.cell(row=row, column=2, value=round(protection['unprotected_size_gb'], 1))
    row += 1
    ws.cell(row=row, column=1, value="Total Unprotected Size (TB)")
    ws.cell(row=row, column=2, value=round(protection['unprotected_size_gb'] / 1024, 2))
    row += 3

    # === Resource List ===
    row = write_section_header(ws, row, "Unprotected Resource List",
                                "(Sorted by size, largest first)")

    headers = [
        "Priority", "Provider", "Account", "Region", "Resource Type",
        "Resource Name", "Resource ID", "Size (GB)", "Environment"
    ]
    write_header_row(ws, row, headers)
    row += 1

    for idx, r in enumerate(unprotected_sorted[:500], 1):  # Limit to 500 rows
        tags = r.get('tags', {}) or {}
        env = tags.get('Environment', tags.get('environment',
               tags.get('Env', tags.get('env', ''))))

        write_data_row(ws, row, [
            idx,
            get_provider(r.get('resource_type', '')),
            r.get('account_id', ''),
            r.get('region', ''),
            r.get('resource_type', ''),
            r.get('name', ''),
            r.get('resource_id', ''),
            r.get('size_gb', 0) or 0,
            env
        ])

        # Highlight production resources
        if env and env.lower() in ['prod', 'production', 'prd']:
            for col in range(1, 10):
                ws.cell(row=row, column=col).fill = STATUS_COLORS['unprotected']

        row += 1

    if len(unprotected_sorted) > 500:
        ws.cell(row=row, column=1, value=f"... and {len(unprotected_sorted) - 500} more resources")

    # Set column widths
    set_column_widths(ws, {
        'A': 10, 'B': 12, 'C': 20, 'D': 15, 'E': 30,
        'F': 30, 'G': 40, 'H': 12, 'I': 15
    })

    ws.freeze_panes = 'A7'
    if unprotected_sorted:
        ws.auto_filter.ref = f"A6:I{row - 1}"


