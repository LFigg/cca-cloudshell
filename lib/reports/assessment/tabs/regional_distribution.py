"""Regional Distribution tab: resources by region for cluster placement planning."""
from typing import Dict, List

from openpyxl import Workbook

from ..analysis import (
    _get_region_group,
    analyze_regions,
)
from ..excel_helpers import (
    set_column_widths,
    write_data_row,
    write_header_row,
    write_section_header,
)
from ..styles import (
    TITLE_FONT,
)


def generate_regional_distribution(wb: Workbook, resources: List[Dict]) -> None:
    """Generate Regional Distribution tab."""
    ws = wb.create_sheet(title="Regional Distribution")

    row = 1

    # Title
    ws.cell(row=row, column=1, value="Regional Distribution").font = TITLE_FONT
    ws.cell(row=row + 1, column=1, value="(Use for Cohesity cluster placement planning)")
    row += 3

    regions = analyze_regions(resources)

    # === Summary by Region ===
    row = write_section_header(ws, row, "Summary by Region")

    write_header_row(ws, row, ["Region", "Provider(s)", "Resource Count", "Size (GB)", "Size (TB)"])
    row += 1

    # Sort by size descending for better visibility
    sorted_regions = sorted(regions.items(), key=lambda x: x[1]['size_gb'], reverse=True)

    for region, data in sorted_regions:
        providers = ', '.join(sorted(data['providers']))
        write_data_row(ws, row, [
            region,
            providers,
            data['count'],
            round(data['size_gb'], 1),
            round(data['size_gb'] / 1024, 2)
        ])
        row += 1

    row += 2

    # === Cluster Placement Recommendations ===
    row = write_section_header(ws, row, "Cluster Placement Recommendations")

    ws.cell(row=row, column=1, value="Grouped by geographic proximity (nearby regions consolidated):")
    row += 1

    # Group regions by geographic area
    region_groups: Dict[str, Dict] = {}

    for region, data in sorted_regions:
        if data['size_gb'] < 1:  # Skip negligible regions
            continue
        group_name, preferred = _get_region_group(region)
        if group_name not in region_groups:
            region_groups[group_name] = {'regions': [], 'size_gb': 0, 'count': 0, 'preferred': preferred}
        region_groups[group_name]['regions'].append(region)
        region_groups[group_name]['size_gb'] += data['size_gb']
        region_groups[group_name]['count'] += data['count']

    # Sort groups by size
    sorted_groups = sorted(region_groups.items(), key=lambda x: x[1]['size_gb'], reverse=True)

    # Filter to significant groups (>100 GB)
    significant_groups = [(name, data) for name, data in sorted_groups if data['size_gb'] > 100]

    if significant_groups:
        write_header_row(ws, row, ["Priority", "Geographic Area", "Regions", "Size (TB)", "Recommendation"])
        row += 1

        for idx, (group_name, data) in enumerate(significant_groups, 1):
            size_tb = data['size_gb'] / 1024
            regions_str = ', '.join(sorted(data['regions']))
            preferred = data['preferred']

            if size_tb > 50:
                rec = f"Primary cluster in {preferred}"
            elif size_tb > 10:
                rec = f"Dedicated cluster in {preferred}"
            else:
                rec = "Can be protected from primary cluster"

            write_data_row(ws, row, [idx, group_name, regions_str, round(size_tb, 2), rec])
            row += 1
    else:
        ws.cell(row=row, column=1, value="No regions with >100 GB of data")

    # Set column widths
    set_column_widths(ws, {'A': 25, 'B': 20, 'C': 15, 'D': 15, 'E': 45})

    ws.freeze_panes = 'A5'


