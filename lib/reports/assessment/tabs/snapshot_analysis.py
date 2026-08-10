"""Snapshot Analysis tab: automated backup pattern breakdown (DLM, scripts, AWS Backup)."""
from typing import Dict, List

from openpyxl import Workbook

from ..analysis import (
    analyze_snapshot_patterns,
)
from ..excel_helpers import (
    set_column_widths,
    write_data_row,
    write_header_row,
    write_section_header,
)
from ..styles import (
    HEADER_FONT,
    TITLE_FONT,
)


def generate_snapshot_analysis(wb: Workbook, resources: List[Dict]) -> None:
    """Generate Snapshot Analysis tab showing automated backup patterns."""
    ws = wb.create_sheet(title="Snapshot Analysis")

    row = 1

    # Title
    ws.cell(row=row, column=1, value="Snapshot Automation Analysis").font = TITLE_FONT
    ws.cell(row=row + 1, column=1, value="(Identifies automated backup sources outside AWS Backup)")
    row += 3

    # Get snapshot analysis
    analysis = analyze_snapshot_patterns(resources)

    if not analysis or not analysis.get('has_data'):
        ws.cell(row=row, column=1, value="No snapshot data found")
        return

    total_ebs = analysis.get('total_ebs_snapshots', 0)
    categories = analysis.get('categories', {})

    # === EBS Snapshot Sources Section ===
    if total_ebs > 0:
        row = write_section_header(ws, row, "EBS Snapshot Sources",
                                   f"Total: {total_ebs:,} snapshots")

        write_header_row(ws, row, ["Source", "Count", "% of Total", "Notes"])
        row += 1

        # Combine script categories for display
        script_total = (categories.get('script_daily', 0) + categories.get('script_weekly', 0) +
                       categories.get('script_monthly', 0) + categories.get('script_other', 0))

        sources = [
            ('AWS Backup', categories.get('aws_backup', 0), 'AWS Backup service'),
            ('Data Lifecycle Manager (DLM)', categories.get('dlm', 0), 'AWS DLM policies'),
            ('Script-based (Daily)', categories.get('script_daily', 0), 'Automated scripts - daily schedule'),
            ('Script-based (Weekly)', categories.get('script_weekly', 0), 'Automated scripts - weekly schedule'),
            ('Script-based (Monthly)', categories.get('script_monthly', 0), 'Automated scripts - monthly schedule'),
            ('Script-based (Other)', categories.get('script_other', 0), 'Automated with hourly/backup/snapshot patterns'),
            ('AMI Artifacts', categories.get('ami_artifact', 0), 'Created alongside AMIs'),
            ('Cross-Region/Account Copies', categories.get('cross_copy', 0), 'Copied from other regions/accounts'),
            ('Manual/Custom', categories.get('manual', 0), 'Other or unidentified'),
        ]

        for name, count, notes in sources:
            if count > 0:
                pct = (count / total_ebs * 100) if total_ebs else 0
                write_data_row(ws, row, [name, count, f"{pct:.1f}%", notes])
                row += 1
        row += 1

        # Summary of non-AWS-Backup automation
        non_awsbackup_automated = categories.get('dlm', 0) + script_total
        if non_awsbackup_automated > 0:
            row = write_section_header(ws, row, "Automated Backups Summary", "")
            ws.cell(row=row, column=1, value="Total automated (non-AWS-Backup):")
            ws.cell(row=row, column=2, value=non_awsbackup_automated)
            pct = (non_awsbackup_automated / total_ebs * 100) if total_ebs else 0
            ws.cell(row=row, column=3, value=f"{pct:.1f}%")
            row += 2

        # === DLM Policies Section ===
        dlm_policies = analysis.get('dlm_policies', {})
        if dlm_policies:
            row = write_section_header(ws, row, "DLM Policy Details",
                                       f"({len(dlm_policies)} policies found)")

            write_header_row(ws, row, ["Policy ID", "Snapshot Count", "Total Size (GB)", "Schedule"])
            row += 1

            # Sort by count descending
            sorted_policies = sorted(dlm_policies.items(),
                                    key=lambda x: x[1]['count'], reverse=True)

            for policy_id, info in sorted_policies[:20]:
                write_data_row(ws, row, [
                    policy_id or 'Unknown',
                    info.get('count', 0),
                    round(info.get('size_gb', 0), 1),
                    info.get('schedule_name', 'Unknown')
                ])
                row += 1
            row += 1

        # === Top Automated Patterns Section ===
        desc_patterns = analysis.get('desc_patterns', {})
        if desc_patterns:
            row = write_section_header(ws, row, "Top Automated Backup Patterns",
                                       "(Description patterns indicating scripted backups)")

            write_header_row(ws, row, ["Pattern/Description", "Count", "% of Total"])
            row += 1

            # Sort by count, show top 15
            sorted_patterns = sorted(desc_patterns.items(),
                                    key=lambda x: x[1], reverse=True)

            for pattern, count in sorted_patterns[:15]:
                if count >= 10:  # Only show significant patterns
                    pct = (count / total_ebs * 100) if total_ebs else 0
                    write_data_row(ws, row, [pattern[:60], count, f"{pct:.1f}%"])
                    row += 1
            row += 1

        # === Schedule Analysis ===
        schedule = analysis.get('schedule_analysis', {})
        if schedule and schedule.get('peak_hours'):
            row = write_section_header(ws, row, "Schedule Patterns",
                                       "(Detected from snapshot creation times)")

            # Peak hours
            ws.cell(row=row, column=1, value="Peak Creation Hours (UTC):").font = HEADER_FONT
            row += 1
            for hour, pct in schedule['peak_hours'][:5]:
                ws.cell(row=row, column=1, value=f"  {hour:02d}:00")
                ws.cell(row=row, column=2, value=f"{pct:.1f}%")
                row += 1
            row += 1

            # Day distribution
            if schedule.get('dow_distribution'):
                ws.cell(row=row, column=1, value="Day of Week Distribution:").font = HEADER_FONT
                row += 1
                for day, pct in schedule['dow_distribution']:
                    ws.cell(row=row, column=1, value=f"  {day}")
                    ws.cell(row=row, column=2, value=f"{pct:.1f}%")
                    row += 1
                row += 1

            # Scheduling indicator
            if schedule.get('likely_scheduled'):
                ws.cell(row=row, column=1, value="⚠ High concentration in specific hours suggests automated scheduling")
                row += 2

        # === Retention Analysis ===
        retention = analysis.get('retention_analysis', {})
        if retention:
            row = write_section_header(ws, row, "Retention Analysis",
                                       "(Age distribution of snapshots)")

            write_header_row(ws, row, ["Age Range", "Count", "% of Total"])
            row += 1

            age_ranges = [
                ('< 7 days', retention.get('under_7_days', 0)),
                ('7-14 days', retention.get('7_to_14_days', 0)),
                ('14-30 days', retention.get('14_to_30_days', 0)),
                ('30-90 days', retention.get('30_to_90_days', 0)),
                ('90-365 days', retention.get('90_to_365_days', 0)),
                ('> 1 year', retention.get('over_365_days', 0)),
            ]

            for range_name, count in age_ranges:
                if count > 0:
                    pct = (count / total_ebs * 100) if total_ebs else 0
                    write_data_row(ws, row, [range_name, count, f"{pct:.1f}%"])
                    row += 1
            row += 1

            # Inferred retention policies
            inferred = retention.get('inferred_policies', [])
            if inferred:
                ws.cell(row=row, column=1, value="Inferred Retention Policies:").font = HEADER_FONT
                row += 1
                for policy in inferred:
                    ws.cell(row=row, column=1, value=f"  • {policy}")
                    row += 1
                row += 1

    # === RDS Snapshot Analysis ===
    rds_analysis = analysis.get('rds_analysis', {})
    total_rds = analysis.get('total_rds_snapshots', 0)
    if total_rds > 0:
        row += 1
        row = write_section_header(ws, row, "RDS Snapshot Sources",
                                   f"Total: {total_rds:,} snapshots")

        write_header_row(ws, row, ["Source", "Count", "% of Total"])
        row += 1

        rds_sources = [
            ('AWS Backup', rds_analysis.get('aws_backup', 0)),
            ('Automated (RDS)', rds_analysis.get('automated', 0)),
            ('Manual', rds_analysis.get('manual', 0)),
            ('Cross-Region Copy', rds_analysis.get('cross_region', 0)),
        ]

        for name, count in rds_sources:
            if count > 0:
                pct = (count / total_rds * 100) if total_rds else 0
                write_data_row(ws, row, [name, count, f"{pct:.1f}%"])
                row += 1

    # Set column widths
    set_column_widths(ws, {'A': 50, 'B': 15, 'C': 12, 'D': 50})

    ws.freeze_panes = 'A3'


