"""Data Quality tab: resources whose actual usage could not be measured (see lib/data_quality.py)."""
from typing import Dict, List

from openpyxl import Workbook

from lib.data_quality import compute_data_quality_summary_from_dicts

from ..excel_helpers import (
    format_percent,
    set_column_widths,
    write_data_row,
    write_header_row,
    write_section_header,
)
from ..styles import (
    STATUS_COLORS,
    TITLE_FONT,
)


def generate_data_quality_tab(wb: Workbook, resources: List[Dict]) -> None:
    """Generate Data Quality tab.

    Surfaces every resource type where actual usage couldn't be measured (real
    usage or nothing - see lib/data_quality.py's policy), so a customer-facing
    reader can see the gap explicitly instead of it being silently invisible
    inside a 0 GB total elsewhere in this report. Always created (even with no
    gaps) so its absence is never mistaken for "this wasn't checked."
    """
    ws = wb.create_sheet(title="Data Quality")

    row = 1
    ws.cell(row=row, column=1, value="Data Quality").font = TITLE_FONT
    ws.cell(
        row=row + 1, column=1,
        value=(
            "Resource types whose actual usage could not be measured are reported as 0 GB here and "
            "throughout this workbook - never an allocated/provisioned/quota estimate standing in for "
            "real usage. This tab makes those gaps visible rather than leaving them indistinguishable "
            "from a genuinely empty resource."
        )
    )
    row += 3

    dq_summary = compute_data_quality_summary_from_dicts(resources)

    if not dq_summary:
        ws.cell(row=row, column=1, value="No resource types in this inventory participate in this policy.")
        set_column_widths(ws, {'A': 90})
        return

    if dq_summary['total_resources_unavailable'] == 0:
        row = write_section_header(ws, row, "No gaps found")
        ws.cell(
            row=row, column=1,
            value="Every tracked resource type in this inventory has real, measured usage data."
        )
        row += 2
    else:
        row = write_section_header(
            ws, row, "Resource types with unmeasured usage",
            f"{dq_summary['total_resources_unavailable']} resource(s) across "
            f"{len(dq_summary['resource_types_with_gaps'])} type(s) - "
            f"~{dq_summary['total_unmeasured_context_gb']:,.1f} GB of allocated/estimated capacity is "
            "known but excluded from every size total in this workbook"
        )

        write_header_row(ws, row, [
            "Resource Type", "Unmeasured", "Total", "% Unmeasured",
            "Allocated/Estimated Capacity (GB, context only - NOT sized)"
        ])
        row += 1

        for bucket in dq_summary['resource_types_with_gaps']:
            pct = (bucket['unavailable'] / bucket['total'] * 100) if bucket['total'] else 0
            write_data_row(ws, row, [
                bucket['resource_type'],
                bucket['unavailable'],
                bucket['total'],
                format_percent(pct),
                round(bucket['unmeasured_context_gb'], 1) if bucket['unmeasured_context_gb'] else '',
            ])
            ws.cell(row=row, column=1).fill = STATUS_COLORS['partial']
            row += 1

        row += 1

    # Full breakdown, including fully-measured resource types, for completeness
    row = write_section_header(ws, row, "All tracked resource types")
    write_header_row(ws, row, ["Resource Type", "Measured", "Unavailable", "Not Applicable", "Total"])
    row += 1
    for bucket in dq_summary['all_tracked_resource_types']:
        write_data_row(ws, row, [
            bucket['resource_type'], bucket['measured'], bucket['unavailable'],
            bucket['not_applicable'], bucket['total'],
        ])
        row += 1

    set_column_widths(ws, {'A': 32, 'B': 14, 'C': 14, 'D': 16, 'E': 45})
    ws.freeze_panes = 'A5'


