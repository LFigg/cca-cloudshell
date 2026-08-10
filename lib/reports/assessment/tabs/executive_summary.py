"""Executive Summary tab: environment overview, sizing summary, protection status."""
from collections import defaultdict
from datetime import datetime, timezone
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Font

from lib.data_quality import compute_data_quality_summary_from_dicts

from ..analysis import (
    analyze_protection_status,
    categorize_resources,
    get_provider,
)
from ..excel_helpers import (
    format_currency,
    format_percent,
    set_column_widths,
    write_data_row,
    write_header_row,
    write_section_header,
)
from ..styles import (
    PROVIDER_COLORS,
    STATUS_COLORS,
    TITLE_FONT,
)


def generate_executive_summary(wb: Workbook, resources: List[Dict],
                                cost_data: Dict, metadata: Dict) -> None:
    """Generate Executive Summary tab."""
    ws = wb.active
    assert ws is not None, "Workbook must have an active sheet"
    ws.title = "Executive Summary"

    row = 1

    # Title
    ws.cell(row=row, column=1, value="Cohesity Assessment Report").font = TITLE_FONT
    row += 1
    ws.cell(row=row, column=1, value=f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}")
    row += 3

    # === Environment Overview Section ===
    row = write_section_header(ws, row, "Environment Overview")

    # Count by provider
    provider_counts = defaultdict(int)
    provider_sizes = defaultdict(float)
    for r in resources:
        provider = get_provider(r.get('resource_type', ''))
        provider_counts[provider] += 1
        provider_sizes[provider] += r.get('size_gb', 0) or 0

    overview_data = [
        ("Total Resources", len(resources)),
        ("Total Orgs/Tenants", len(metadata.get('orgs', []))),
        ("Total Accounts/Subscriptions", len(metadata.get('accounts', []))),
        ("Cloud Providers", ', '.join(sorted(provider_counts.keys()))),
    ]

    for label, value in overview_data:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1

    # Provider breakdown
    write_header_row(ws, row, ["Provider", "Resources", "Size (GB)"])
    row += 1

    for provider in sorted(provider_counts.keys()):
        write_data_row(ws, row, [
            provider,
            provider_counts[provider],
            round(provider_sizes[provider], 1)
        ])
        # Color code provider
        if provider in PROVIDER_COLORS:
            ws.cell(row=row, column=1).fill = PROVIDER_COLORS[provider]
            ws.cell(row=row, column=1).font = Font(color="FFFFFF", bold=True)
        row += 1

    row += 2

    # === Sizing Summary Section ===
    row = write_section_header(ws, row, "Sizing Summary",
                                "(For Cohesity sizing calculator)")

    categories = categorize_resources(resources)

    write_header_row(ws, row, ["Workload Category", "Count", "Size (GB)", "Size (TB)"])
    row += 1

    total_size = 0
    for category in sorted(categories.keys()):
        data = categories[category]
        if data['count'] > 0 and category not in ['Snapshots', 'Backup Services', 'Other']:
            size_gb = data['size_gb']
            total_size += size_gb
            write_data_row(ws, row, [
                category,
                data['count'],
                round(size_gb, 1),
                round(size_gb / 1024, 2)
            ])
            row += 1

    # Total row
    ws.cell(row=row, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=row, column=3, value=round(total_size, 1)).font = Font(bold=True)
    ws.cell(row=row, column=4, value=round(total_size / 1024, 2)).font = Font(bold=True)
    row += 1

    # Data-quality caveat: some resource types in the total above may be
    # reporting 0 GB not because they're empty, but because their actual usage
    # couldn't be measured (size_source='unavailable', see lib/data_quality.py) -
    # this total is a floor, not a guarantee, whenever that gap is non-empty.
    dq_summary = compute_data_quality_summary_from_dicts(resources)
    if dq_summary and dq_summary['total_resources_unavailable'] > 0:
        note = (
            f"⚠ TOTAL may understate real capacity: {dq_summary['total_resources_unavailable']} "
            f"resource(s) across {len(dq_summary['resource_types_with_gaps'])} type(s) have no measured "
            "actual usage (reported as 0 GB above, not an allocated estimate) - see the Data Quality tab"
        )
        ws.cell(row=row, column=1, value=note).fill = STATUS_COLORS['partial']
        row += 1

    row += 1

    # === Protection Status Section ===
    row = write_section_header(ws, row, "Protection Status")

    protection = analyze_protection_status(resources)

    prot_data = [
        ("Total Protectable Resources", protection['total_protectable']),
        ("Protected", protection['protected_count']),
        ("Unprotected", protection['unprotected_count']),
        ("Coverage", format_percent(protection['coverage_percent'])),
    ]

    for label, value in prot_data:
        ws.cell(row=row, column=1, value=label)
        cell = ws.cell(row=row, column=2, value=value)
        # Color code protection status
        if label == "Protected":
            cell.fill = STATUS_COLORS['protected']
        elif label == "Unprotected":
            cell.fill = STATUS_COLORS['unprotected']
        row += 1

    row += 1

    # Size breakdown
    write_header_row(ws, row, ["Status", "Size (GB)", "Size (TB)"])
    row += 1

    for status, size_key, fill in [
        ("Protected", 'protected_size_gb', STATUS_COLORS['protected']),
        ("Unprotected", 'unprotected_size_gb', STATUS_COLORS['unprotected']),
    ]:
        size_gb = protection[size_key]
        write_data_row(ws, row, [status, round(size_gb, 1), round(size_gb / 1024, 2)])
        ws.cell(row=row, column=1).fill = fill
        row += 1

    row += 2

    # === Current Backup Cost Section ===
    if cost_data.get('total_cost', 0) > 0:
        row = write_section_header(ws, row, "Current Backup/Storage Costs",
                                    "(Monthly snapshot and backup storage costs)")

        ws.cell(row=row, column=1, value="Total Monthly Cost")
        ws.cell(row=row, column=2, value=format_currency(cost_data['total_cost']))
        row += 2

        if cost_data.get('by_provider'):
            write_header_row(ws, row, ["Provider", "Monthly Cost"])
            row += 1

            for provider, pdata in sorted(cost_data['by_provider'].items()):
                write_data_row(ws, row, [
                    provider.upper(),
                    format_currency(pdata.get('total', 0))
                ])
                row += 1

    # Set column widths
    set_column_widths(ws, {'A': 30, 'B': 20, 'C': 15, 'D': 15})

