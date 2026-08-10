"""Sizing Inputs tab: workload inventory by type for the Cohesity sizing calculator."""
from collections import defaultdict
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font

from ..analysis import (
    _get_db_engine_group,
    _get_region_group,
    _is_kubernetes_node,
    analyze_protection_status,
    categorize_resources,
    get_provider,
    get_workload_category,
)
from ..excel_helpers import (
    set_column_widths,
    write_data_row,
    write_header_row,
    write_section_header,
)
from ..styles import (
    TITLE_FONT,
)


def generate_sizing_inputs(wb: Workbook, resources: List[Dict],
                          change_rate_data: Optional[Dict[str, Any]] = None) -> None:
    """Generate Sizing Inputs tab with complete coverage, protected-only, regional, and encryption breakdown."""
    ws = wb.create_sheet(title="Sizing Inputs")

    row = 1

    # Title
    ws.cell(row=row, column=1, value="Cohesity Sizing Calculator Inputs").font = TITLE_FONT
    row += 2

    # Categorize all resources
    categorize_resources(resources)

    # Get protection status for "apples to apples" sizing
    protection = analyze_protection_status(resources)
    protected_resources = protection.get('protected_resources', [])

    # Default change rates by category (including database-specific rates)
    change_rates = {
        'Virtual Machines': 3.0,
        'Block Storage': 2.0,
        'Block Storage (Unattached)': 2.0,
        'File Storage': 1.0,
        'Object Storage': 0.5,
        'Kubernetes/Containers': 3.0,
        'Cache/In-Memory': 2.0,
        # Database-specific change rates (transaction log generation)
        'DB: MySQL/MariaDB': 5.0,
        'DB: PostgreSQL': 5.0,
        'DB: SQL Server': 5.0,
        'DB: Oracle': 5.0,
        'DB: Cosmos DB': 3.0,
        'DB: DocumentDB': 4.0,
        'DB: Neptune': 4.0,
        'DB: DynamoDB': 2.0,
        'DB: Redshift': 3.0,
        'DB: Synapse': 3.0,
        'DB: BigTable': 2.0,
        'DB: Spanner': 3.0,
        'Databases': 5.0,  # Fallback for unrecognized
    }

    priority_order = ['Virtual Machines', 'Block Storage', 'Block Storage (Unattached)',
                      'DB: MySQL/MariaDB', 'DB: PostgreSQL', 'DB: SQL Server', 'DB: Oracle',
                      'DB: Cosmos DB', 'DB: DocumentDB', 'DB: Neptune', 'DB: DynamoDB',
                      'DB: Redshift', 'DB: Synapse', 'DB: BigTable', 'DB: Spanner',
                      'File Storage', 'Object Storage', 'Kubernetes/Containers', 'Cache/In-Memory']

    # ==========================================================================
    # SECTION 1: Complete Coverage (Full Environment)
    # ==========================================================================
    row = write_section_header(ws, row, "Option 1: Complete Coverage",
                                "(Full environment protection - input into Cohesity sizing calculator)")

    # Re-categorize with database breakdown
    categories_with_db: Dict[str, Dict[str, Any]] = {}

    # Build volume lookup for VM storage calculation
    volume_by_id: Dict[str, Dict] = {}
    for r in resources:
        if r.get('resource_type') == 'aws:ec2:volume':
            volume_by_id[r.get('resource_id', '')] = r

    attached_volume_ids: set = set()

    for r in resources:
        rtype = r.get('resource_type', '')
        meta = r.get('metadata', {}) or {}
        size_gb = r.get('size_gb', 0) or 0

        # Skip replicas and snapshots
        if meta.get('is_read_replica'):
            continue
        if 'snapshot' in rtype.lower():
            continue

        # Skip DataWarehouse tier Azure SQL databases - they're Synapse dedicated SQL pools
        # and are counted separately as azure:synapse:sqlpool
        if rtype == 'azure:sql:database' and meta.get('tier') == 'DataWarehouse':
            continue

        # Check if database - get specific engine
        db_engine = _get_db_engine_group(r)
        if db_engine:
            category = db_engine
        elif rtype == 'aws:ec2:instance':
            # Calculate storage from attached volumes
            attached_vols = meta.get('attached_volumes', [])
            size_gb = 0
            for vol_id in attached_vols:
                if vol_id in volume_by_id:
                    size_gb += volume_by_id[vol_id].get('size_gb', 0) or 0
                    attached_volume_ids.add(vol_id)
            if _is_kubernetes_node(r):
                category = 'Kubernetes/Containers'
            else:
                category = 'Virtual Machines'
        elif rtype in ['azure:vm', 'gcp:compute:instance']:
            if _is_kubernetes_node(r):
                category = 'Kubernetes/Containers'
            else:
                category = 'Virtual Machines'
        elif rtype == 'aws:ec2:volume':
            if r.get('resource_id', '') in attached_volume_ids:
                continue  # Already counted with VM
            attached = meta.get('attached_to')
            if attached:
                continue
            category = 'Block Storage (Unattached)'
        elif rtype in ['azure:disk', 'gcp:compute:disk']:
            attached = meta.get('attached_to')
            if attached:
                continue
            category = 'Block Storage (Unattached)'
        elif rtype in ['aws:efs:filesystem', 'aws:fsx:filesystem', 'azure:storage:fileshare',
                       'gcp:filestore:instance']:
            category = 'File Storage'
        elif rtype in ['aws:s3:bucket', 'azure:storage:blob', 'gcp:storage:bucket']:
            category = 'Object Storage'
        elif rtype in ['aws:eks:cluster', 'azure:aks:cluster', 'gcp:container:cluster']:
            category = 'Kubernetes/Containers'
        elif rtype in ['aws:elasticache:cluster', 'azure:redis:cache', 'gcp:redis:instance']:
            category = 'Cache/In-Memory'
        else:
            cat = get_workload_category(rtype)
            if cat in ['Snapshots', 'Backup Services', 'Other']:
                continue
            category = cat

        if category not in categories_with_db:
            categories_with_db[category] = {'count': 0, 'size_gb': 0}
        categories_with_db[category]['count'] += 1
        categories_with_db[category]['size_gb'] += size_gb

    write_header_row(ws, row, [
        "Workload Type", "Count", "Size (GB)", "Size (TB)",
        "Daily Change Rate (%)", "Est. Daily Change (GB)"
    ])
    row += 1

    total_size = 0
    total_change = 0

    sorted_categories = sorted(
        categories_with_db.keys(),
        key=lambda x: priority_order.index(x) if x in priority_order else 100
    )

    for category in sorted_categories:
        data = categories_with_db[category]
        if category in ['Snapshots', 'Backup Services', 'Other']:
            continue
        if data['count'] == 0:
            continue

        size_gb = data['size_gb']
        total_size += size_gb

        change_rate = change_rates.get(category, 2.0)
        daily_change = size_gb * (change_rate / 100)
        total_change += daily_change

        write_data_row(ws, row, [
            category,
            data['count'],
            round(size_gb, 1),
            round(size_gb / 1024, 2),
            change_rate,
            round(daily_change, 1)
        ])
        row += 1

    # Total row
    ws.cell(row=row, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=row, column=3, value=round(total_size, 1)).font = Font(bold=True)
    ws.cell(row=row, column=4, value=round(total_size / 1024, 2)).font = Font(bold=True)
    ws.cell(row=row, column=6, value=round(total_change, 1)).font = Font(bold=True)
    row += 3

    # ==========================================================================
    # SECTION 2: Currently Protected Only (Apples to Apples)
    # ==========================================================================
    row = write_section_header(ws, row, "Option 2: Currently Protected Only (Apples-to-Apples)",
                                "(Migrate existing backup coverage to Cohesity - same scope)")

    # Categorize only protected resources (reuse volume_by_id from above)
    protected_categories: Dict[str, Dict] = {}
    protected_volume_ids: set = set()

    for r in protected_resources:
        rtype = r.get('resource_type', '')
        meta = r.get('metadata', {}) or {}

        # Check if database - get specific engine
        db_engine = _get_db_engine_group(r)
        if db_engine:
            category = db_engine
            size_gb = r.get('size_gb', 0) or 0

        elif rtype == 'aws:ec2:instance':
            if _is_kubernetes_node(r):
                category = 'Kubernetes/Containers'
            else:
                category = 'Virtual Machines'

            # Calculate storage from attached volumes
            attached_vols = meta.get('attached_volumes', [])
            size_gb = 0
            for vol_id in attached_vols:
                if vol_id in volume_by_id:
                    size_gb += volume_by_id[vol_id].get('size_gb', 0) or 0
                    protected_volume_ids.add(vol_id)

        elif rtype == 'aws:ec2:volume':
            # Check if already counted via an instance
            if r.get('resource_id', '') in protected_volume_ids:
                continue
            category = 'Block Storage (Unattached)'
            size_gb = r.get('size_gb', 0) or 0

        elif rtype in ['aws:efs:filesystem', 'aws:fsx:filesystem']:
            category = 'File Storage'
            size_gb = r.get('size_gb', 0) or 0

        elif rtype == 'aws:s3:bucket':
            category = 'Object Storage'
            size_gb = r.get('size_gb', 0) or 0

        elif rtype in ['aws:eks:cluster', 'azure:aks:cluster', 'gcp:container:cluster']:
            category = 'Kubernetes/Containers'
            size_gb = r.get('size_gb', 0) or 0

        else:
            # Generic handling for other types
            category = get_workload_category(rtype)
            if category in ['Snapshots', 'Backup Services', 'Other']:
                continue
            size_gb = r.get('size_gb', 0) or 0

        if category not in protected_categories:
            protected_categories[category] = {'count': 0, 'size_gb': 0}
        protected_categories[category]['count'] += 1
        protected_categories[category]['size_gb'] += size_gb

    write_header_row(ws, row, [
        "Workload Type", "Count", "Size (GB)", "Size (TB)",
        "Daily Change Rate (%)", "Est. Daily Change (GB)"
    ])
    row += 1

    protected_total_size = 0
    protected_total_change = 0

    sorted_protected = sorted(
        protected_categories.keys(),
        key=lambda x: priority_order.index(x) if x in priority_order else 100
    )

    for category in sorted_protected:
        data = protected_categories[category]
        if data['count'] == 0:
            continue

        size_gb = data['size_gb']
        protected_total_size += size_gb

        change_rate = change_rates.get(category, 2.0)
        daily_change = size_gb * (change_rate / 100)
        protected_total_change += daily_change

        write_data_row(ws, row, [
            category,
            data['count'],
            round(size_gb, 1),
            round(size_gb / 1024, 2),
            change_rate,
            round(daily_change, 1)
        ])
        row += 1

    # Total row for protected
    ws.cell(row=row, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=row, column=3, value=round(protected_total_size, 1)).font = Font(bold=True)
    ws.cell(row=row, column=4, value=round(protected_total_size / 1024, 2)).font = Font(bold=True)
    ws.cell(row=row, column=6, value=round(protected_total_change, 1)).font = Font(bold=True)
    row += 2

    # Coverage comparison
    if total_size > 0:
        coverage_pct = (protected_total_size / total_size) * 100
        ws.cell(row=row, column=1, value="Currently Protected Coverage:")
        ws.cell(row=row, column=2, value=f"{coverage_pct:.1f}% of total environment")
        row += 1

        gap_size = total_size - protected_total_size
        ws.cell(row=row, column=1, value="Unprotected Gap:")
        ws.cell(row=row, column=2, value=f"{gap_size/1024:.1f} TB ({100-coverage_pct:.1f}%)")
    row += 3

    # ==========================================================================
    # SECTION 3: Regional Breakdown for CE Cluster Planning
    # ==========================================================================
    row = write_section_header(ws, row, "Option 3: Regional Breakdown (Per-Region CE Clusters)",
                                "(Each region group typically requires a dedicated Cohesity Cloud Edition cluster)")

    # Build volume lookup for VM storage calculation (reuse if available)
    regional_volume_by_id: Dict[str, Dict] = {}
    for r in resources:
        if r.get('resource_type') == 'aws:ec2:volume':
            regional_volume_by_id[r.get('resource_id', '')] = r

    # Track volumes attached to VMs (to avoid double-counting)
    regional_attached_vol_ids: set = set()

    # Build regional breakdown with encryption status
    regional_data: Dict[str, Dict[str, Dict[str, Any]]] = {}  # region_group -> category -> {count, size, encrypted_size, ...}

    for r in resources:
        rtype = r.get('resource_type', '')
        meta = r.get('metadata', {}) or {}
        region = r.get('region', 'unknown')

        # Skip replicas and snapshots
        if meta.get('is_read_replica'):
            continue
        if 'snapshot' in rtype.lower():
            continue

        # Get region group
        region_group, _ = _get_region_group(region)

        # Check if database - get specific engine first
        db_engine = _get_db_engine_group(r)
        if db_engine:
            category = db_engine
            size_gb = r.get('size_gb', 0) or 0
        elif rtype == 'aws:ec2:instance':
            if _is_kubernetes_node(r):
                category = 'Kubernetes/Containers'
            else:
                category = 'Virtual Machines'
            # Calculate storage from attached volumes
            size_gb = 0
            attached_vols = meta.get('attached_volumes', [])
            for vol_id in attached_vols:
                if vol_id in regional_volume_by_id:
                    size_gb += regional_volume_by_id[vol_id].get('size_gb', 0) or 0
                    regional_attached_vol_ids.add(vol_id)
        elif rtype == 'aws:ec2:volume':
            # Skip volumes attached to instances (already counted)
            vol_id = r.get('resource_id', '')
            if vol_id in regional_attached_vol_ids:
                continue
            attached = meta.get('attached_instance')
            if attached:
                continue  # Skip attached volumes
            category = 'Block Storage (Unattached)'
            size_gb = r.get('size_gb', 0) or 0
        elif rtype in ['azure:vm', 'gcp:compute:instance']:
            if _is_kubernetes_node(r):
                category = 'Kubernetes/Containers'
            else:
                category = 'Virtual Machines'
            # Azure/GCP may have size_gb directly or in attached disks
            size_gb = r.get('size_gb', 0) or 0
        elif rtype in ['azure:disk', 'gcp:compute:disk']:
            attached = meta.get('attached_to')
            if attached:
                continue
            category = 'Block Storage (Unattached)'
            size_gb = r.get('size_gb', 0) or 0
        elif rtype in ['aws:efs:filesystem', 'aws:fsx:filesystem', 'azure:storage:fileshare',
                       'gcp:filestore:instance']:
            category = 'File Storage'
            size_gb = r.get('size_gb', 0) or 0
        elif rtype in ['aws:s3:bucket', 'azure:storage:blob', 'gcp:storage:bucket']:
            category = 'Object Storage'
            size_gb = r.get('size_gb', 0) or 0
        elif rtype in ['aws:eks:cluster', 'azure:aks:cluster', 'gcp:container:cluster']:
            category = 'Kubernetes/Containers'
            size_gb = r.get('size_gb', 0) or 0
        else:
            cat = get_workload_category(rtype)
            if cat in ['Snapshots', 'Backup Services', 'Other']:
                continue
            category = cat
            size_gb = r.get('size_gb', 0) or 0

        # Check encryption status - TDE and guest-level encryption affect dedupe/compression
        # Server-side encryption (Azure EncryptionAtRestWithPlatformKey, AWS KMS, GCP default)
        # is transparent and does NOT affect Cohesity dedupe
        is_encrypted = False

        # Check for TDE on database resources
        # Azure SQL has TDE enabled by default since 2017
        if rtype in ['azure:sql:database', 'azure:sql:managedinstance']:
            is_encrypted = meta.get('tde_enabled', True)  # Default True for Azure SQL
        elif meta.get('tde_enabled', False):
            is_encrypted = True
        elif rtype in ['aws:ec2:volume', 'aws:rds:instance', 'aws:rds:cluster',
                     'aws:efs:filesystem', 'aws:fsx:filesystem']:
            # AWS: Check for guest-level encryption indicators
            # Server-side KMS encryption (encrypted=True) does NOT affect dedupe
            # Only guest-level (LUKS, BitLocker) would affect dedupe - not detectable from API
            is_encrypted = False  # AWS server-side encryption is transparent
        elif rtype.startswith('azure:'):
            # Azure: Server-side encryption (EncryptionAtRestWithPlatformKey/CustomerKey) is transparent
            # Only Azure Disk Encryption (ADE) with BitLocker/dm-crypt affects dedupe
            enc_type = meta.get('encryption_type', '')
            # ADE would show different encryption_type, platform keys are server-side
            is_encrypted = enc_type and enc_type not in [
                'EncryptionAtRestWithPlatformKey',
                'EncryptionAtRestWithCustomerKey',
                'EncryptionAtRestWithPlatformAndCustomerKeys',
                ''
            ]
        elif rtype.startswith('gcp:'):
            # GCP: Default encryption is server-side and transparent
            # Only CSEK (Customer-Supplied Encryption Keys) or guest-level affects dedupe
            is_encrypted = meta.get('guest_os_encrypted', False)

        # Initialize region group if needed
        if region_group not in regional_data:
            regional_data[region_group] = {}
        if category not in regional_data[region_group]:
            regional_data[region_group][category] = {
                'count': 0, 'size_gb': 0,
                'encrypted_count': 0, 'encrypted_size_gb': 0,
                'unencrypted_count': 0, 'unencrypted_size_gb': 0
            }

        # Update stats
        regional_data[region_group][category]['count'] += 1
        regional_data[region_group][category]['size_gb'] += size_gb
        if is_encrypted:
            regional_data[region_group][category]['encrypted_count'] += 1
            regional_data[region_group][category]['encrypted_size_gb'] += size_gb
        else:
            regional_data[region_group][category]['unencrypted_count'] += 1
            regional_data[region_group][category]['unencrypted_size_gb'] += size_gb

    # Sort region groups by total size
    sorted_regions = sorted(
        regional_data.keys(),
        key=lambda rg: sum(d['size_gb'] for d in regional_data[rg].values()),
        reverse=True
    )

    # Output each region group as a sub-section
    for region_group in sorted_regions:
        region_total_size = sum(d['size_gb'] for d in regional_data[region_group].values())
        if region_total_size < 1:  # Skip negligible regions
            continue

        # Region header
        ws.cell(row=row, column=1, value=f"Region: {region_group}").font = Font(bold=True, color="0000AA")
        ws.cell(row=row, column=2, value=f"Total: {region_total_size/1024:.2f} TB")
        row += 1

        write_header_row(ws, row, [
            "Workload Type", "Count", "Size (TB)",
            "Encrypted (TB)", "Unencrypted (TB)", "Daily Change Rate (%)"
        ])
        row += 1

        region_categories = regional_data[region_group]
        sorted_cats = sorted(
            region_categories.keys(),
            key=lambda x: priority_order.index(x) if x in priority_order else 100
        )

        region_total = 0
        for category in sorted_cats:
            data = region_categories[category]
            if data['count'] == 0:
                continue

            size_tb = data['size_gb'] / 1024
            enc_tb = data['encrypted_size_gb'] / 1024
            unenc_tb = data['unencrypted_size_gb'] / 1024
            cr = change_rates.get(category, 2.0)
            region_total += data['size_gb']

            write_data_row(ws, row, [
                category,
                data['count'],
                round(size_tb, 2),
                round(enc_tb, 2),
                round(unenc_tb, 2),
                cr
            ])
            row += 1

        # Region total
        ws.cell(row=row, column=1, value="Region Total").font = Font(bold=True)
        ws.cell(row=row, column=3, value=round(region_total / 1024, 2)).font = Font(bold=True)
        row += 2

    row += 1

    # ==========================================================================
    # SECTION 4: Workloads by Encryption Status
    # ==========================================================================
    row = write_section_header(ws, row, "Option 4: Workloads by Encryption Status",
                                "(Encrypted workloads achieve 30-50% less deduplication)")

    # Build volume lookup for VM storage calculation (same as Section 3)
    enc_volume_by_id: Dict[str, Dict] = {}
    for r in resources:
        if r.get('resource_type') == 'aws:ec2:volume':
            enc_volume_by_id[r.get('resource_id', '')] = r

    # Track volumes attached to VMs
    enc_attached_vol_ids: set = set()

    # Build encryption-separated workload summary
    enc_workloads: Dict[str, Dict[str, Any]] = {}  # "Category - Encrypted/Unencrypted" -> stats

    for r in resources:
        rtype = r.get('resource_type', '')
        meta = r.get('metadata', {}) or {}

        if meta.get('is_read_replica'):
            continue
        if 'snapshot' in rtype.lower():
            continue

        # Calculate size_gb (special handling for VMs)
        size_gb = 0
        is_encrypted = False  # Only guest-level encryption or TDE affects dedupe

        # Check if database - get specific engine first
        db_engine = _get_db_engine_group(r)
        if db_engine:
            category = db_engine
            size_gb = r.get('size_gb', 0) or 0
            # TDE (Transparent Data Encryption) DOES affect dedupe/compression
            # Azure SQL has TDE enabled by default since 2017
            if rtype in ['azure:sql:database', 'azure:sql:managedinstance']:
                is_encrypted = meta.get('tde_enabled', True)  # Default True for Azure SQL
            else:
                is_encrypted = meta.get('tde_enabled', False)
        elif rtype == 'aws:ec2:instance':
            # Calculate storage from attached volumes
            attached_vols = meta.get('attached_volumes', [])
            for vol_id in attached_vols:
                if vol_id in enc_volume_by_id:
                    vol = enc_volume_by_id[vol_id]
                    vol_size = vol.get('size_gb', 0) or 0
                    size_gb += vol_size
                    enc_attached_vol_ids.add(vol_id)
                    # AWS EBS encryption is server-side (KMS) - transparent, doesn't affect dedupe
            if _is_kubernetes_node(r):
                category = 'Kubernetes/Containers'
            else:
                category = 'Virtual Machines'
        elif rtype == 'aws:ec2:volume':
            # Skip volumes attached to instances
            vol_id = r.get('resource_id', '')
            if vol_id in enc_attached_vol_ids:
                continue
            if meta.get('attached_instance'):
                continue
            category = 'Block Storage'
            size_gb = r.get('size_gb', 0) or 0
            # AWS EBS encryption is server-side (KMS) - transparent, doesn't affect dedupe
            is_encrypted = False
        else:
            # Determine category
            category = get_workload_category(rtype)
            if category in ['Snapshots', 'Backup Services', 'Other']:
                continue
            size_gb = r.get('size_gb', 0) or 0
            # Server-side encryption (AWS KMS, Azure platform keys, GCP default) doesn't affect dedupe
            # Only guest-level encryption would affect dedupe - not detectable from cloud API
            is_encrypted = False

        # Create key with encryption status
        enc_status = "Encrypted" if is_encrypted else "Unencrypted"
        key = f"{category} - {enc_status}"

        if key not in enc_workloads:
            enc_workloads[key] = {'count': 0, 'size_gb': 0, 'category': category, 'encrypted': is_encrypted}

        enc_workloads[key]['count'] += 1
        enc_workloads[key]['size_gb'] += size_gb

    write_header_row(ws, row, [
        "Workload Type", "Encryption", "Count", "Size (GB)", "Size (TB)",
        "Daily Change Rate (%)", "Est. Daily Change (GB)"
    ])
    row += 1

    # Sort: by category first, then encrypted last (show unencrypted first)
    sorted_enc = sorted(
        enc_workloads.keys(),
        key=lambda k: (
            priority_order.index(enc_workloads[k]['category']) if enc_workloads[k]['category'] in priority_order else 100,
            1 if enc_workloads[k]['encrypted'] else 0
        )
    )

    enc_total_size = 0
    enc_total_change = 0

    for key in sorted_enc:
        data = enc_workloads[key]
        if data['count'] == 0:
            continue

        size_gb = data['size_gb']
        enc_total_size += size_gb

        cr = change_rates.get(data['category'], 2.0)
        daily_change = size_gb * (cr / 100)
        enc_total_change += daily_change

        enc_label = "Yes" if data['encrypted'] else "No"

        write_data_row(ws, row, [
            data['category'],
            enc_label,
            data['count'],
            round(size_gb, 1),
            round(size_gb / 1024, 2),
            cr,
            round(daily_change, 1)
        ])
        row += 1

    # Total
    ws.cell(row=row, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=row, column=4, value=round(enc_total_size, 1)).font = Font(bold=True)
    ws.cell(row=row, column=5, value=round(enc_total_size / 1024, 2)).font = Font(bold=True)
    ws.cell(row=row, column=7, value=round(enc_total_change, 1)).font = Font(bold=True)
    row += 2

    # Sizing note
    ws.cell(row=row, column=1, value="Note: Use encrypted workloads with reduced dedupe estimates in sizing calculator.")
    row += 3

    # ==========================================================================
    # SECTION 5: Encryption Summary (affects deduplication efficiency)
    # ==========================================================================
    row = write_section_header(ws, row, "Encryption Summary",
                                "(Overall encryption stats for sizing deduplication estimates)")

    # Analyze encryption across all resources
    encrypted_data = {'count': 0, 'size_gb': 0}
    unencrypted_data = {'count': 0, 'size_gb': 0}
    encryption_by_type: Dict[str, Dict[str, Any]] = {}

    for r in resources:
        rtype = r.get('resource_type', '')
        meta = r.get('metadata', {}) or {}
        size_gb = r.get('size_gb', 0) or 0

        # Skip replicas and snapshots
        if meta.get('is_read_replica'):
            continue
        if 'snapshot' in rtype.lower():
            continue

        # Check encryption status based on resource type
        # Server-side encryption (AWS KMS, Azure platform keys, GCP default) is transparent
        # TDE (Transparent Data Encryption) DOES affect dedupe - check tde_enabled
        # Guest-level encryption (BitLocker, LUKS) would affect dedupe but is not detectable
        is_encrypted = False

        # Check for TDE on database resources
        # Azure SQL has TDE enabled by default since 2017
        if rtype in ['azure:sql:database', 'azure:sql:managedinstance']:
            is_encrypted = meta.get('tde_enabled', True)  # Default True for Azure SQL
        elif meta.get('tde_enabled', False):
            is_encrypted = True

        if is_encrypted:
            encrypted_data['count'] += 1
            encrypted_data['size_gb'] += size_gb
        else:
            unencrypted_data['count'] += 1
            unencrypted_data['size_gb'] += size_gb

        # Track by type
        if rtype not in encryption_by_type:
            encryption_by_type[rtype] = {'encrypted': 0, 'encrypted_gb': 0,
                                         'unencrypted': 0, 'unencrypted_gb': 0}
        if is_encrypted:
            encryption_by_type[rtype]['encrypted'] += 1
            encryption_by_type[rtype]['encrypted_gb'] += size_gb
        else:
            encryption_by_type[rtype]['unencrypted'] += 1
            encryption_by_type[rtype]['unencrypted_gb'] += size_gb

    # Summary
    total_analyzed = encrypted_data['size_gb'] + unencrypted_data['size_gb']
    encrypted_pct = (encrypted_data['size_gb'] / total_analyzed * 100) if total_analyzed > 0 else 0

    write_header_row(ws, row, ["Status", "Resource Count", "Size (GB)", "Size (TB)", "% of Total"])
    row += 1

    write_data_row(ws, row, [
        "Encrypted",
        encrypted_data['count'],
        round(encrypted_data['size_gb'], 1),
        round(encrypted_data['size_gb'] / 1024, 2),
        f"{encrypted_pct:.1f}%"
    ])
    row += 1

    write_data_row(ws, row, [
        "Unencrypted",
        unencrypted_data['count'],
        round(unencrypted_data['size_gb'], 1),
        round(unencrypted_data['size_gb'] / 1024, 2),
        f"{100-encrypted_pct:.1f}%"
    ])
    row += 2

    # Note about sizing implications
    ws.cell(row=row, column=1, value="Note: Encrypted data typically achieves 30-50% less deduplication.")
    row += 1
    ws.cell(row=row, column=1, value="Consider this when sizing Cohesity storage capacity.")
    row += 3

    # ==========================================================================
    # SECTION 6: Database Sizing Details (Transaction Logs vs Data)
    # ==========================================================================
    row = write_section_header(ws, row, "Database Sizing Details",
                                "(Transaction logs require separate sizing from data)")

    # Collect database details by engine
    db_types = ['aws:rds:instance', 'aws:rds:cluster', 'azure:sql:database',
                'azure:sql:managedinstance', 'azure:cosmosdb:account', 'gcp:sql:instance']

    db_by_engine: Dict[str, Dict[str, Any]] = {}

    for r in resources:
        rtype = r.get('resource_type', '')
        if rtype not in db_types:
            continue

        meta = r.get('metadata', {}) or {}

        # Skip read replicas
        if meta.get('is_read_replica'):
            continue

        size_gb = r.get('size_gb', 0) or 0

        # Determine engine type
        engine = 'Unknown'
        if rtype in ['aws:rds:instance', 'aws:rds:cluster']:
            engine = meta.get('engine', 'Unknown')
        elif rtype == 'azure:sql:database':
            engine = 'SQL Server (Azure)'
        elif rtype == 'azure:sql:managedinstance':
            engine = 'SQL Server (Azure MI)'
        elif rtype == 'azure:cosmosdb:account':
            engine = 'Cosmos DB'
        elif rtype == 'gcp:sql:instance':
            db_version = meta.get('database_version', '')
            if 'MYSQL' in db_version.upper():
                engine = 'MySQL (GCP)'
            elif 'POSTGRES' in db_version.upper():
                engine = 'PostgreSQL (GCP)'
            elif 'SQLSERVER' in db_version.upper():
                engine = 'SQL Server (GCP)'
            else:
                engine = db_version or 'Unknown (GCP)'

        # Normalize similar engines
        engine_lower = engine.lower()
        if 'mysql' in engine_lower or 'mariadb' in engine_lower or 'aurora-mysql' in engine_lower:
            engine_group = 'MySQL/MariaDB'
        elif 'postgres' in engine_lower or 'aurora-postgresql' in engine_lower:
            engine_group = 'PostgreSQL'
        elif 'sqlserver' in engine_lower or 'sql server' in engine_lower:
            engine_group = 'SQL Server'
        elif 'oracle' in engine_lower:
            engine_group = 'Oracle'
        elif 'cosmos' in engine_lower:
            engine_group = 'NoSQL (Cosmos DB)'
        elif 'docdb' in engine_lower or 'documentdb' in engine_lower:
            engine_group = 'NoSQL (DocumentDB)'
        elif 'neptune' in engine_lower:
            engine_group = 'Graph DB (Neptune)'
        elif 'dynamodb' in engine_lower:
            engine_group = 'NoSQL (DynamoDB)'
        else:
            engine_group = engine

        if engine_group not in db_by_engine:
            db_by_engine[engine_group] = {
                'count': 0,
                'data_size_gb': 0,
                'encrypted_count': 0,
            }

        db_by_engine[engine_group]['count'] += 1
        db_by_engine[engine_group]['data_size_gb'] += size_gb
        # TDE (Transparent Data Encryption) DOES affect dedupe/compression
        # Azure SQL has TDE enabled by default since 2017
        if rtype in ['azure:sql:database', 'azure:sql:managedinstance']:
            if meta.get('tde_enabled', True):  # Default True for Azure SQL
                db_by_engine[engine_group]['encrypted_count'] += 1
        elif meta.get('tde_enabled', False):
            db_by_engine[engine_group]['encrypted_count'] += 1

    if db_by_engine:
        # Typical transaction log generation rates by engine (as % of data size per day)
        # These are industry estimates for OLTP workloads - used as fallback when no actual data
        tlog_rate_estimates = {
            'MySQL/MariaDB': 10.0,         # Binary logs, row-based replication
            'PostgreSQL': 15.0,            # WAL logs, MVCC overhead
            'SQL Server': 12.0,            # Transaction log, full recovery
            'Oracle': 10.0,                # Redo logs
            'NoSQL (Cosmos DB)': 5.0,      # Change feed
            'NoSQL (DocumentDB)': 8.0,     # Oplog (MongoDB compatible)
            'NoSQL (DynamoDB)': 3.0,       # Streams
            'Graph DB (Neptune)': 8.0,     # Journal logs
        }

        # Map engine groups to change rate data keys
        engine_to_cr_keys = {
            'MySQL/MariaDB': ['aws:rds-mysql', 'aws:rds-mariadb', 'aws:rds-aurora-mysql',
                             'gcp:cloudsql-mysql', 'azure:mysql'],
            'PostgreSQL': ['aws:rds-postgres', 'aws:rds-aurora-postgresql',
                          'gcp:cloudsql-postgres', 'azure:postgres'],
            'SQL Server': ['aws:rds-sqlserver', 'gcp:cloudsql-sqlserver',
                          'azure:sql-database', 'azure:sql-managedinstance'],
            'Oracle': ['aws:rds-oracle'],
            'NoSQL (Cosmos DB)': ['azure:cosmosdb'],
            'NoSQL (DocumentDB)': ['aws:documentdb'],
            'NoSQL (DynamoDB)': ['aws:dynamodb'],
            'Graph DB (Neptune)': ['aws:neptune'],
        }

        # Check if we have actual change rate data
        has_actual_data = change_rate_data and change_rate_data.get('has_actual_data', False)
        actual_change_rates = change_rate_data.get('change_rates', {}) if change_rate_data and has_actual_data else {}

        # Function to get actual tlog rate for an engine group
        def get_actual_tlog_gb(engine_group: str, data_size_gb: float) -> Tuple[Optional[float], bool]:
            """Returns (daily_tlog_gb, is_actual) tuple."""
            if not actual_change_rates:
                return None, False

            cr_keys = engine_to_cr_keys.get(engine_group, [])
            total_actual_tlog = 0.0
            found_actual = False

            for key in cr_keys:
                if key in actual_change_rates:
                    cr = actual_change_rates[key]
                    tlog = cr.get('transaction_logs')
                    if tlog and 'daily_generation_gb' in tlog:
                        total_actual_tlog += tlog.get('daily_generation_gb', 0)
                        found_actual = True

            if found_actual:
                return total_actual_tlog, True
            return None, False

        write_header_row(ws, row, [
            "Database Engine", "Count", "Data Size (TB)",
            "Daily Tlog (GB)", "Monthly Tlog (TB)", "Source", "Notes"
        ])
        row += 1

        total_db_size = 0
        total_tlog_daily = 0
        any_actual = False

        for engine_group in sorted(db_by_engine.keys()):
            data = db_by_engine[engine_group]
            data_size_tb = data['data_size_gb'] / 1024
            total_db_size += data['data_size_gb']

            # Try to get actual tlog data, fall back to estimate
            actual_tlog, is_actual = get_actual_tlog_gb(engine_group, data['data_size_gb'])

            if is_actual and actual_tlog is not None:
                daily_tlog_gb = actual_tlog
                source = "Actual"
                any_actual = True
            else:
                # Estimate transaction log generation
                tlog_rate = tlog_rate_estimates.get(engine_group, 8.0)  # Default 8%
                daily_tlog_gb = data['data_size_gb'] * (tlog_rate / 100)
                source = "Estimated"

            monthly_tlog_tb = (daily_tlog_gb * 30) / 1024
            total_tlog_daily += daily_tlog_gb

            # Encryption note
            enc_pct = (data['encrypted_count'] / data['count'] * 100) if data['count'] > 0 else 0
            notes = f"{enc_pct:.0f}% encrypted"
            if enc_pct > 50:
                notes += " (reduced dedupe)"

            write_data_row(ws, row, [
                engine_group,
                data['count'],
                round(data_size_tb, 2),
                round(daily_tlog_gb, 1),
                round(monthly_tlog_tb, 2),
                source,
                notes
            ])
            row += 1

        # Totals
        ws.cell(row=row, column=1, value="TOTAL").font = Font(bold=True)
        ws.cell(row=row, column=2, value=sum(d['count'] for d in db_by_engine.values())).font = Font(bold=True)
        ws.cell(row=row, column=3, value=round(total_db_size / 1024, 2)).font = Font(bold=True)
        ws.cell(row=row, column=4, value=round(total_tlog_daily, 1)).font = Font(bold=True)
        ws.cell(row=row, column=5, value=round((total_tlog_daily * 30) / 1024, 2)).font = Font(bold=True)
        row += 2

        # Sizing guidance
        ws.cell(row=row, column=1, value="Database Backup Sizing Notes:")
        row += 1
        ws.cell(row=row, column=1, value="• Data backups: Use incremental forever (changed blocks only)")
        row += 1
        ws.cell(row=row, column=1, value="• Transaction logs: Require 100% capture (full log shipping)")
        row += 1
        if any_actual:
            ws.cell(row=row, column=1, value="• Tlog rates marked 'Actual' are from CloudWatch metrics (7-day average)")
            row += 1
            ws.cell(row=row, column=1, value="• Tlog rates marked 'Estimated' use industry-standard assumptions")
        else:
            ws.cell(row=row, column=1, value="• Tlog rates above are estimates - actual rates depend on workload activity")
            row += 1
            ws.cell(row=row, column=1, value="• For actual tlog sizing, re-run collection without --skip-change-rate flag")
        row += 1
    else:
        ws.cell(row=row, column=1, value="No databases found in inventory")
        row += 1

    row += 2

    # === Detailed Breakdown Section ===
    row = write_section_header(ws, row, "Detailed Resource Breakdown",
                                "(Resource counts by type)")

    # Count by resource type
    type_counts = defaultdict(lambda: {'count': 0, 'size_gb': 0})
    for r in resources:
        rtype = r.get('resource_type', 'unknown')
        type_counts[rtype]['count'] += 1
        type_counts[rtype]['size_gb'] += r.get('size_gb', 0) or 0

    write_header_row(ws, row, ["Resource Type", "Provider", "Count", "Size (GB)"])
    row += 1

    for rtype in sorted(type_counts.keys()):
        data = type_counts[rtype]
        write_data_row(ws, row, [
            rtype,
            get_provider(rtype),
            data['count'],
            round(data['size_gb'], 1)
        ])
        row += 1

    # Set column widths
    set_column_widths(ws, {'A': 35, 'B': 15, 'C': 12, 'D': 15, 'E': 20, 'F': 20})

    # Freeze header
    ws.freeze_panes = 'A4'


