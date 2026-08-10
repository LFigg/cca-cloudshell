"""Generic openpyxl cell/formatting helpers shared across every report tab."""
from typing import Any, Dict, List, Optional

from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from .styles import HEADER_FILL, HEADER_FONT, SECTION_FONT, THIN_BORDER


def set_cell(ws: Worksheet, row: int, col: int, value: Any,
             font: Optional[Font] = None, fill: Optional[PatternFill] = None,
             border: Optional[Border] = None, alignment: Optional[Alignment] = None) -> None:
    """Set cell value with optional styling."""
    cell = ws.cell(row=row, column=col, value=value)
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if border:
        cell.border = border
    if alignment:
        cell.alignment = alignment


def write_header_row(ws: Worksheet, row: int, headers: List[str], start_col: int = 1) -> None:
    """Write a styled header row."""
    for col_idx, header in enumerate(headers, start=start_col):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER


def write_data_row(ws: Worksheet, row: int, data: List[Any], start_col: int = 1,
                   fill: Optional[PatternFill] = None) -> None:
    """Write a data row with optional styling."""
    for col_idx, value in enumerate(data, start=start_col):
        cell = ws.cell(row=row, column=col_idx, value=value)
        cell.border = THIN_BORDER
        if fill:
            cell.fill = fill


def write_section_header(ws: Worksheet, row: int, title: str,
                         subtitle: Optional[str] = None) -> int:
    """Write a section header, return next row."""
    ws.cell(row=row, column=1, value=title).font = SECTION_FONT
    if subtitle:
        ws.cell(row=row + 1, column=1, value=subtitle)
        return row + 3
    return row + 2


def set_column_widths(ws: Worksheet, widths: Dict[str, int]) -> None:
    """Set column widths from a dict mapping column letters to widths."""
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def format_number(value: float, decimals: int = 1) -> str:
    """Format number with thousands separator."""
    if value is None:
        return ''
    if decimals == 0:
        return f"{int(value):,}"
    return f"{value:,.{decimals}f}"


def format_percent(value: float) -> str:
    """Format percentage."""
    if value is None:
        return ''
    return f"{value:.1f}%"


def format_currency(value: float, currency: str = 'USD') -> str:
    """Format currency value."""
    if value is None:
        return ''
    if currency == 'USD':
        return f"${value:,.2f}"
    return f"{value:,.2f} {currency}"
